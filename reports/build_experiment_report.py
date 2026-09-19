"""Build the illustrated submission report from measured artifacts."""
from pathlib import Path
import json
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage
R=Path(__file__).resolve().parents[1]; O=R/'reports'; A=R/'artifacts/experiment'
s=json.loads((A/'summary.json').read_text());qa=json.loads((O/'screenshots/qa.json').read_text())
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyCustom',fontName='Helvetica',fontSize=10.5,leading=15,textColor=colors.HexColor('#34303C'),spaceAfter=10))
styles.add(ParagraphStyle(name='TitleCustom',fontName='Helvetica-Bold',fontSize=29,leading=34,textColor=colors.HexColor('#38173e'),spaceAfter=18))
styles.add(ParagraphStyle(name='CaptionCustom',fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor('#65546b'),spaceBefore=8,spaceAfter=14))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Caption','Heading 1','Heading 2']:
 doc.styles[name].font.name='Calibri';doc.styles[name].font.size=Pt(10.5 if name=='Normal' else 9 if name=='Caption' else 23 if name=='Heading 1' else 15)
 doc.styles[name].font.color.rgb=RGBColor.from_string('38173E' if 'Heading' in name else '34303C')
footer=sec.footer.paragraphs[0];footer.text='ProotyPie | CNN classification study                                        '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
story=[]
def text(t):
 story.append(Paragraph(escape(t),styles['BodyCustom']));doc.add_paragraph(t)
def heading(t):
 story.append(Paragraph(escape(t),styles['TitleCustom']));doc.add_heading(t,1)
def page(title):
 if story:story.append(PageBreak());doc.add_page_break()
 heading(title)
def image(p,caption,maxh=320):
 p=Path(p);im=PILImage.open(p);
 if p.name.startswith('result-'):
  im=im.crop((288,0,1440,850));p=p.with_name('detail-'+p.name);im.save(p)
 w,h=im.size;scale=min(480/w,maxh/h);story.append(Image(str(p),width=w*scale,height=h*scale));story.append(Paragraph(escape(caption),styles['CaptionCustom']));doc.add_picture(str(p),width=Inches(w*scale/72));doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths=None):
 cells=[[Paragraph(escape(str(c)),styles['CaptionCustom']) for c in row]for row in [headers]+rows]
 t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eee5f0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8f6f9')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#935073'))]));story.append(t)
 dt=doc.add_table(rows=1,cols=len(headers));dt.style='Light Shading Accent 1'
 for c,v in zip(dt.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(dt.add_row().cells,row):c.text=str(v)
page('CNN image classification\nActivation & learning-rate study')
text('Machine-learning mini-project | Fruits-360')
text('Abstract. A compact convolutional neural network was trained using four hidden activation functions and three SGD learning rates. All twelve configurations used identical data splits, architecture, initialization seed and training budget. Accuracy, macro precision, macro recall and macro F1 were compared on validation data. The selected configuration was refitted on the development set and evaluated on a held-out official-Test subset.')
text('Dataset scope: Apple Braeburn, Banana, Orange and Strawberry only; 320 training, 80 validation and 160 test images. This is a four-class subset study, not a full Fruits-360 benchmark.')
text('Objective: compare ReLU, Sigmoid, Tanh and Leaky ReLU; investigate learning-rate effects on convergence; identify the strongest tested configuration using validation performance.')
image(A/'validation_comparison.png','Measured validation macro F1 across all twelve configurations.',260)
text('Selected configuration: Tanh with SGD learning rate 0.1. Conclusions apply only to this dataset subset and six-epoch budget.')
page('Dataset & experimental controls')
text('Source: Fruits-360, Horea94/Fruit-Images-Dataset on GitHub. The experiment uses Apple Braeburn, Banana, Orange and Strawberry. Images are converted to RGB, resized to 48 × 48 pixels and scaled to [0, 1].')
table(['Split','Per class','Total','Role'],[['Training',80,320,'Gradient updates'],['Validation',20,80,'Configuration selection'],['Official Test subset',40,160,'Final evaluation']],[105,65,60,250])
text('Sampling is deterministic: the first 100 lexicographically sorted Training JPGs and first 40 Test JPGs per class. Training images are stratified into an 80/20 train/validation split with seed 20260919. SHA-256 checks reject identical files across splits.')
text('Controls: identical architecture, initialization seed, batch size 16 and six epochs for every comparison. Only the hidden activation and learning rate vary. The output always uses softmax. The selected configuration is refitted on all 400 development images.')
text('Important qualification: exact-file hashing does not establish independence between visually similar rotational views. This is a narrow, clean-background experiment, not the complete Fruits-360 benchmark. Random seeds were fixed, but repeated-seed uncertainty was not measured.')
page('Architecture & optimization')
table(['Stage','Operation / output'],[['Input','48 × 48 × 3 RGB image'],['Feature extraction 1','Conv2D: 8 filters, 3 × 3, same padding; chosen activation'],['Downsampling','MaxPooling2D: reduces spatial dimensions'],['Feature extraction 2','Conv2D: 16 filters, 3 × 3; same chosen activation'],['Aggregation','GlobalAveragePooling2D: one value per feature map'],['Decision','Dense: four softmax probabilities']],[135,345])
text('Training objective: sparse categorical cross-entropy. Mini-batch SGD updates each parameter in the direction opposite the loss gradient: new weight = old weight − learning rate × gradient. A larger learning rate takes bigger steps, but can also make optimization unstable.')
text('ReLU keeps positive values and zeros negative values. Sigmoid compresses values into (0, 1). Tanh maps values into (−1, 1). Leaky ReLU retains a small negative slope (0.1). These choices affect the signals and gradients flowing through the convolutional layers.')
text('Accuracy is the fraction of correct predictions. Precision = TP / (TP + FP); recall = TP / (TP + FN); F1 = 2 × precision × recall / (precision + recall). Macro scores are the unweighted mean of per-class scores; undefined precision is assigned zero.')
page('Twelve configurations. One comparison.')
rows=[]
for e in s['experiments']:
 m=e['validation_metrics'];rows.append([e['activation'],str(e['learning_rate'])]+[f'{m[k]:.4f}' for k in ['accuracy','precision_macro','recall_macro','f1_macro']])
table(['Activation','LR','Accuracy','Precision*','Recall*','F1*'],rows,[100,48,80,84,84,84])
text('* Macro averaging gives every fruit class equal weight. These are validation metrics, not twelve separate test-set evaluations.')
text('Tanh with SGD learning rate 0.1 had the highest validation macro F1 (1.0000). The lowest learning rate left all four activations at 25% validation accuracy after six epochs. That is evidence of inadequate learning within this budget—not proof that those configurations could never converge.')
image(A/'validation_comparison.png','Validation-only model comparison. Higher is better; all runs share the same data split and epoch budget.',210)
page('Learning behaviour')
image(A/'tanh_lr0p1_loss.png','Selected configuration: training and validation cross-entropy over six epochs.',235)
image(A/'tanh_lr0p1_accuracy.png','Selected configuration: training and validation accuracy. Epoch indices on these saved plots begin at zero.',235)
text('The selected run ends with training loss 0.4681 and validation loss 0.4476. Higher learning rates performed better within the tested short budget. Six epochs are not enough to establish asymptotic convergence; no universal best activation is claimed.')
page('Held-out evaluation')
image(A/'test_confusion_matrix.png','Confusion matrix for the final refitted model on 160 official-Test subset images: 40 images in each class.',330)
table(['Metric','Measured value'],[[k.replace('_',' '),f'{v:.4f}'] for k,v in s['test_metrics'].items()],[240,240])
text('All 160 selected test images were correctly classified. This is a real result on a small, visually constrained subset—not a promise of perfect performance on unseen fruit varieties, backgrounds or phone-camera images. The model is closed-set and will assign even unknown objects to one of its four classes.')
page('Learning-rate effects across activations')
text('At learning rate 0.001, every activation remained at 25% validation accuracy after six epochs. At 0.1, validation accuracy reached 75% for ReLU, 48.75% for Sigmoid, 100% for Tanh and 77.5% for Leaky ReLU. This supports faster useful progress at the higher tested rate within the fixed budget, not a universal optimal learning rate.')
image(A/'relu_lr0p1_loss.png','ReLU, learning rate 0.1: training and validation loss.',230)
image(A/'sigmoid_lr0p1_loss.png','Sigmoid, learning rate 0.1: loss remains comparatively high.',230)
page('Convergence discussion & conclusion')
image(A/'leaky_relu_lr0p1_loss.png','Leaky ReLU, learning rate 0.1: measured loss trajectory.',250)
text('Tanh at 0.1 produced the highest validation macro F1 (1.0000), followed by Leaky ReLU at 0.1 (0.7064) and ReLU at 0.1 (0.6540). At low learning rates, limited movement within six epochs should be described as slow learning under this budget, not proof of inability to learn.')
text('The selected configuration achieved 1.0000 accuracy, macro precision, macro recall and macro F1 on the 160-image official-Test subset after refitting on 400 development images. Perfect subset performance does not establish real-world generalization.')
text('Limitations: only four classes; lexicographic sampling; one fixed seed; six epochs; clean backgrounds; possible similarity between rotational views despite no identical-file overlap. Neither asymptotic convergence nor statistical superiority across repeated seeds was established.')
text('Conclusion. Activation function and learning rate substantially affected short-budget training. Tanh with SGD at 0.1 was the best of the twelve tested configurations. A broader dataset and repeated runs are needed before generalizing this ranking.')
page('Reproduction & references')
text('Run python experiments/run_fruits360_experiment.py --run from the repository root after installing requirements.txt. The notebook notebooks/Fruits360_CNN_Experiments.ipynb invokes the same executable implementation. The saved model and machine-readable metrics are in artifacts/experiment/.')
text('All 24 per-configuration training/validation loss and accuracy plots are included in the accompanying experiment-results archive. The notebook can display the comparison chart and confusion matrix. The report contains selected representative curves; the archive retains every configuration.')
text('Dataset citation: Horea Muresan and Mihai Oltean, Fruit recognition from images using deep learning, Acta Universitatis Sapientiae, Informatica, 10(1), 26–42, 2018.')
text('Dataset snapshot: https://github.com/Horea94/Fruit-Images-Dataset (version 2020.05.18.0 according to readme.md). This repository lists 131 classes and has since moved development to https://github.com/fruits-360/. Only the explicitly documented subset was used in this study.')
text('Evidence: experiments/run_fruits360_experiment.py; artifacts/experiment/summary.json; class_names.json; test_confusion_matrix.png; per-configuration learning-curve images. No full-dataset experiment is claimed.')
def decorate(c,d):
 c.setStrokeColor(colors.HexColor('#935073'));c.line(48,803,547,803);c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#65546b'));c.drawString(48,817,'CNN image classification / Fruits-360');c.drawString(48,25,'Fruits-360 subset • Fruit-type classification');c.drawRightString(547,25,str(d.page))
SimpleDocTemplate(str(O/'Fruits360_CNN_Submission_Report.pdf'),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=58,bottomMargin=45,title='Fruits-360 CNN activation and learning-rate study',author='ProotyPie project').build(story,onFirstPage=decorate,onLaterPages=decorate)
doc.save(O/'Fruits360_CNN_Submission_Report.docx');print('Built illustrated PDF and DOCX')
