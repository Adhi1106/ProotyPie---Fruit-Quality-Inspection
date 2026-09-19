"""Build the illustrated submission report from measured artifacts."""
from pathlib import Path
import json
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage
R=Path(__file__).resolve().parents[1]; O=R/'reports'; A=R/'artifacts/experiment'
s=json.loads((A/'summary.json').read_text());qa=json.loads((O/'screenshots/qa.json').read_text())
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyCustom',fontName='Helvetica',fontSize=10.5,leading=15,textColor=colors.HexColor('#34303C'),spaceAfter=10))
styles.add(ParagraphStyle(name='TitleCustom',fontName='Helvetica-Bold',fontSize=29,leading=34,textColor=colors.HexColor('#38173e'),spaceAfter=18))
styles.add(ParagraphStyle(name='CaptionCustom',fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor('#65546b'),spaceBefore=8,spaceAfter=14))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Caption','Heading 1','Heading 2']:
 doc.styles[name].font.name='Calibri';doc.styles[name].font.size=Pt(10.5 if name=='Normal' else 9 if name=='Caption' else 23 if name=='Heading 1' else 15)
 doc.styles[name].font.color.rgb=RGBColor.from_string('38173E' if 'Heading' in name else '34303C')
footer=sec.footer.paragraphs[0];footer.text='ProotyPie | CNN classification study                                        '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
story=[]
def text(t):
 story.append(Paragraph(escape(t),styles['BodyCustom']));doc.add_paragraph(t)
def heading(t):
 story.append(Paragraph(escape(t),styles['TitleCustom']));doc.add_heading(t,1)
def page(title):
 if story:story.append(PageBreak());doc.add_page_break()
 heading(title)
def image(p,caption,maxh=320):
 p=Path(p);im=PILImage.open(p);
 if p.name.startswith('result-'):
  im=im.crop((288,0,1440,850));p=p.with_name('detail-'+p.name);im.save(p)
 w,h=im.size;scale=min(480/w,maxh/h);story.append(Image(str(p),width=w*scale,height=h*scale));story.append(Paragraph(escape(caption),styles['CaptionCustom']));doc.add_picture(str(p),width=Inches(w*scale/72));doc.add_paragraph(caption,'Caption')
def table(headers,rows,widths=None):
 cells=[[Paragraph(escape(str(c)),styles['CaptionCustom']) for c in row]for row in [headers]+rows]
 t=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eee5f0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8f6f9')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#935073'))]));story.append(t)
 dt=doc.add_table(rows=1,cols=len(headers));dt.style='Light Shading Accent 1'
 for c,v in zip(dt.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(dt.add_row().cells,row):c.text=str(v)
page('ProotyPie\nFruit classification, examined.')
text('Machine-learning mini-project • Illustrated submission report')
text('A controlled comparison of activation functions and SGD learning rates, demonstrated through a working fruit-classification application.')
image(O/'screenshots/result-Banana.png','Working application capture: an uploaded Fruits-360 banana classified by the saved CNN. No prediction shown here is mocked.',330)
text('Four fruit classes. Four activation functions. Three learning rates. Twelve measured experiments.')
text('Scope: a documented Fruits-360 subset. This study classifies fruit type; it does not determine freshness or food safety.')
page('Project at a glance')
text('Abstract. This project investigates how activation functions and learning rates affect a compact convolutional neural network. Twelve configurations were trained on the same four-class Fruits-360 subset. Configuration selection used validation macro F1; the selected configuration was refitted on training plus validation images before official-Test subset evaluation. A Next.js interface calls a Python API that loads the trained Keras model.')
table(['Assignment requirement','Evidence in this submission'],[['CNN on a public dataset','Two convolutional layers; four Fruits-360 classes'],['Four activation functions','ReLU, Sigmoid, Tanh and Leaky ReLU'],['Gradient descent and learning rates','Mini-batch SGD at 0.001, 0.01 and 0.1'],['Accuracy, precision, recall and F1','All 12 validation configurations; selected model test metrics'],['Loss, accuracy and confusion plots','24 learning curves, comparison chart and test confusion matrix'],['Application demonstration','Real upload-to-result screenshots and JSON export verification']],[180,300])
text('Reading guide: dataset and controls; architecture and optimization; comparison results; learning curves; test evaluation; application screenshots; verification and limitations; reproduction and references.')
page('Dataset & experimental controls')
text('Source: Fruits-360, Horea94/Fruit-Images-Dataset on GitHub. The experiment uses Apple Braeburn, Banana, Orange and Strawberry. Images are converted to RGB, resized to 48 × 48 pixels and scaled to [0, 1].')
table(['Split','Per class','Total','Role'],[['Training',80,320,'Gradient updates'],['Validation',20,80,'Configuration selection'],['Official Test subset',40,160,'Final evaluation']],[105,65,60,250])
text('Sampling is deterministic: the first 100 lexicographically sorted Training JPGs and first 40 Test JPGs per class. Training images are stratified into an 80/20 train/validation split with seed 20260919. SHA-256 checks reject identical files across splits.')
text('Controls: identical architecture, initialization seed, batch size 16 and six epochs for every comparison. Only the hidden activation and learning rate vary. The output always uses softmax. The selected configuration is refitted on all 400 development images.')
text('Important qualification: exact-file hashing does not establish independence between visually similar rotational views. This is a narrow, clean-background experiment, not the complete Fruits-360 benchmark. Random seeds were fixed, but repeated-seed uncertainty was not measured.')
page('Architecture & optimization')
table(['Stage','Operation / output'],[['Input','48 × 48 × 3 RGB image'],['Feature extraction 1','Conv2D: 8 filters, 3 × 3, same padding; chosen activation'],['Downsampling','MaxPooling2D: reduces spatial dimensions'],['Feature extraction 2','Conv2D: 16 filters, 3 × 3; same chosen activation'],['Aggregation','GlobalAveragePooling2D: one value per feature map'],['Decision','Dense: four softmax probabilities']],[135,345])
text('Training objective: sparse categorical cross-entropy. Mini-batch SGD updates each parameter in the direction opposite the loss gradient: new weight = old weight − learning rate × gradient. A larger learning rate takes bigger steps, but can also make optimization unstable.')
text('ReLU keeps positive values and zeros negative values. Sigmoid compresses values into (0, 1). Tanh maps values into (−1, 1). Leaky ReLU retains a small negative slope (0.1). These choices affect the signals and gradients flowing through the convolutional layers.')
text('Application flow: browser upload → same-origin Next.js API proxy → Python image validation → saved Keras CNN → class probabilities → result card / JSON download. Gemini is not required for this classification path.')
page('Twelve configurations. One comparison.')
rows=[]
for e in s['experiments']:
 m=e['validation_metrics'];rows.append([e['activation'],str(e['learning_rate'])]+[f'{m[k]:.4f}' for k in ['accuracy','precision_macro','recall_macro','f1_macro']])
table(['Activation','LR','Accuracy','Precision*','Recall*','F1*'],rows,[100,48,80,84,84,84])
text('* Macro averaging gives every fruit class equal weight. These are validation metrics, not twelve separate test-set evaluations.')
text('Tanh with SGD learning rate 0.1 had the highest validation macro F1 (1.0000). The lowest learning rate left all four activations at 25% validation accuracy after six epochs. That is evidence of inadequate learning within this budget—not proof that those configurations could never converge.')
image(A/'validation_comparison.png','Validation-only model comparison. Higher is better; all runs share the same data split and epoch budget.',210)
page('Learning behaviour')
image(A/'tanh_lr0p1_loss.png','Selected configuration: training and validation cross-entropy over six epochs.',235)
image(A/'tanh_lr0p1_accuracy.png','Selected configuration: training and validation accuracy. Epoch indices on these saved plots begin at zero.',235)
text('The selected run ends with training loss 0.4681 and validation loss 0.4476. Higher learning rates performed better within the tested short budget. Six epochs are not enough to establish asymptotic convergence; no universal best activation is claimed.')
page('Held-out evaluation')
image(A/'test_confusion_matrix.png','Confusion matrix for the final refitted model on 160 official-Test subset images: 40 images in each class.',330)
table(['Metric','Measured value'],[[k.replace('_',' '),f'{v:.4f}'] for k,v in s['test_metrics'].items()],[240,240])
text('All 160 selected test images were correctly classified. This is a real result on a small, visually constrained subset—not a promise of perfect performance on unseen fruit varieties, backgrounds or phone-camera images. The model is closed-set and will assign even unknown objects to one of its four classes.')
page('The working application')
image(O/'screenshots/01-dashboard.png','Figure: production-build dashboard, with working fruit-classifier readiness status.',335)
text('The interface keeps image upload, preview, classification, model status and result export together. The fruit-classification lab is selected by default. The legacy freshness path is separate and explicitly unavailable without its original model and configuration.')
text('A port mismatch in the API proxy was corrected: the frontend now forwards to the actual default backend at 127.0.0.1:8000. Screenshots were captured from the running production build, not from a design mockup.')
page('Results in the interface')
for fruit in ['Banana','Orange']:
 row=next(r for r in qa['results'] if r['input']==fruit)
 image(O/f'screenshots/result-{fruit}.png',f"{fruit}: real upload {row['file']}; predicted {row['response']['predicted_class']}; model confidence {row['response']['confidence']:.2f}%.",255)
text('Confidence is the model output for one image. It is neither aggregate test accuracy nor a food-safety score.')
page('Mobile & export verification')
image(O/'screenshots/mobile-result.png','Mobile result at a 390-pixel viewport. No horizontal overflow was detected.',410)
text('The Strawberry upload returned a real classification response. The Download JSON button was exercised and its file saved as reports/screenshots/export-example.json. Browser evidence and endpoint responses are retained in reports/screenshots/qa.json.')
page('Verification & limitations')
table(['Check','Outcome'],[['Backend and experiment unit tests','13 passed in the previous combined run; rechecked for this revision'],['Frontend production build','Passed after proxy / readiness changes'],['Browser upload-to-result','Banana, Orange and Strawberry returned valid CNN results'],['JSON export','Downloaded successfully through the real button'],['Mobile layout','390-pixel viewport; no horizontal overflow']],[190,290])
text('The original fresh/rotten model is missing. This report makes no freshness, shelf-life or food-safety performance claim. Gemini-backed features were not tested and are not required for the demonstrated classification flow.')
text('Study limitations: four classes only; deterministic rather than representative sampling; one random seed; short training budget; possible similarity among rotational images; no external-camera test set; no unknown-object rejection; no significance test.')
text('Conclusion. The comparison demonstrates the required CNN workflow and measurable effects of activation and learning-rate choice. Tanh with SGD at 0.1 was best under this particular setup. Extending classes, training duration, seeds and external-image evaluation would strengthen generalization evidence, but is not claimed as completed here.')
page('Reproduce & present')
text('Use Python 3.11 or 3.12 (the legacy API uses cgi, removed in Python 3.13). Install dependencies from requirements.txt. Start the backend with python api/server.py. In frontend, run npm ci, npm run build and npm start. The frontend default proxy targets port 8000.')
text('To rerun all experiments: python experiments/run_fruits360_experiment.py --run. The notebook notebooks/Fruits360_CNN_Experiments.ipynb invokes the same training pipeline. Results, model and class order are saved under artifacts/experiment/.')
text('Suggested demonstration: open the classification lab; upload a held-out Banana image; run classification; explain the confidence-versus-accuracy distinction; download JSON; show the validation comparison and confusion matrix; close with the subset and freshness limitations.')
text('Primary sources and evidence')
for t in ['Dataset: https://github.com/Horea94/Fruit-Images-Dataset','Source splits: /Training and /Test in that repository','Training implementation: experiments/run_fruits360_experiment.py','Measured results: artifacts/experiment/summary.json','Browser captures: reports/screenshots/','Project: https://github.com/Adhi1106/ProotyPie---Fruit-Quality-Inspection'] :text(t)
text('Original repository author: Adhira Praveen, as credited in the original README. No student register number, institution, supervisor signature or experimental outcome has been invented.')
def decorate(c,d):
 c.setStrokeColor(colors.HexColor('#935073'));c.line(48,803,547,803);c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#65546b'));c.drawString(48,817,'ProotyPie / Machine-learning mini-project');c.drawString(48,25,'Fruits-360 subset • Fruit-type classification');c.drawRightString(547,25,str(d.page))
SimpleDocTemplate(str(O/'ProotyPie_Project_Report.pdf'),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=58,bottomMargin=45,title='ProotyPie — CNN classification study',author='ProotyPie project').build(story,onFirstPage=decorate,onLaterPages=decorate)
doc.save(O/'ProotyPie_Project_Report.docx');print('Built illustrated PDF and DOCX')
